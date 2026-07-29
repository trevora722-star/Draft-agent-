#!/usr/bin/env python3
"""
coParenter Compliance Tracker
=============================

Local, agentic legal-discovery pipeline:

  1. Extracts text from a raw coParenter app export ("inputs/coparenter_log.pdf"),
     tagging every line with its source page ("Log Page N").
  2. Splits the log into token-safe chunks (~12k tokens each) so arbitrarily
     large exports never blow past LLM context or timeout limits.
  3. Sends each chunk — together with the Court Order rulebook
     ("inputs/court_order.txt") — to Claude with schema-enforced structured
     outputs, acting as a strictly objective family-law discovery analyst.
  4. Aggregates every extracted event into a single pandas DataFrame and
     deploys a formatted, lawyer-ready spreadsheet to "outputs/".

Usage:
    pip install -r requirements.txt
    export ANTHROPIC_API_KEY=sk-ant-...   # or `ant auth login`
    python main.py
    python main.py --pdf inputs/coparenter_log.pdf --order inputs/court_order.txt

The output spreadsheet contains exactly these columns, in order:
    Timestamp | Event Type | Actor | Court Order Requirement | Actual Event |
    Compliance Status | Fact-Based Justification | Evidence Reference
"""

from __future__ import annotations

import argparse
import random
import re
import sys
import time
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import List, Literal, Optional

import pandas as pd
from pydantic import BaseModel, Field

try:
    from dateutil import parser as dateutil_parser
except ImportError:  # python-dateutil is in requirements.txt, but degrade gracefully
    dateutil_parser = None

import anthropic

# --------------------------------------------------------------------------- #
# Configuration
# --------------------------------------------------------------------------- #

BASE_DIR = Path(__file__).resolve().parent
INPUTS_DIR = BASE_DIR / "inputs"
OUTPUTS_DIR = BASE_DIR / "outputs"

DEFAULT_PDF_PATH = INPUTS_DIR / "coparenter_log.pdf"
DEFAULT_ORDER_PATH = INPUTS_DIR / "court_order.txt"

MODEL = "claude-opus-5"

# Token-safe chunking. The API-side tokenizer averages ~4 characters per token
# for English prose, so 12,000 tokens ~= 48,000 characters. That lands inside
# the requested 10k-15k token window with headroom for the system prompt,
# the rulebook, and the structured JSON response.
CHUNK_TOKEN_TARGET = 12_000
CHARS_PER_TOKEN_ESTIMATE = 4
CHUNK_CHAR_LIMIT = CHUNK_TOKEN_TARGET * CHARS_PER_TOKEN_ESTIMATE

MAX_OUTPUT_TOKENS = 16_000
MAX_API_RETRIES = 4

COMPLIANCE_STATUSES = (
    "Compliant",
    "Violation: Late/No-Show",
    "Violation: Response Delay",
    "Violation: Non-Payment",
    "Violation: Gatekeeping",
    "Unclassified/Neutral",
)

OUTPUT_COLUMNS = [
    "Timestamp",
    "Event Type",
    "Actor",
    "Court Order Requirement",
    "Actual Event",
    "Compliance Status",
    "Fact-Based Justification",
    "Evidence Reference",
]

SYSTEM_PROMPT_TEMPLATE = """\
You are a strictly objective family law discovery expert preparing evidence \
exhibits for court. You analyze raw co-parenting app logs against a controlling \
Court Order and classify each discrete event's compliance status.

NON-NEGOTIABLE ANALYTICAL RULES:

1. TONE — Cold, factual, chronological. Strip ALL emotional, subjective, or \
dramatic vocabulary from your descriptions. Never use words such as \
"hostilely", "refused", "rudely", "angrily", "ignored", "demanded", \
"aggressively", "sadly", or any characterization of tone or intent. State only \
observable facts: who, what, when, where. \
Example — instead of "Father rudely refused the exchange", write \
"Father stated he would not attend the 17:00 exchange."

2. RULEBOOK — The Court Order below is the sole source of truth for what is \
required. Test every event only against its written terms. Do not invent \
requirements that are not in the order, and do not assume standard practices.

3. CLASSIFICATION — Assign each event exactly one status from this closed set:
   - "Compliant"                  (conduct met the order's written requirement)
   - "Violation: Late/No-Show"    (missed or late custody exchange, visit, or deadline for physical presence)
   - "Violation: Response Delay"  (communication answered outside the order's required response window, or never)
   - "Violation: Non-Payment"     (expense, reimbursement, or support obligation not paid as ordered)
   - "Violation: Gatekeeping"     (withholding access, information, or communication contrary to the order)
   - "Unclassified/Neutral"       (routine event the order does not govern, or insufficient data to test)
   If an event is ambiguous or the log data is incomplete, prefer \
"Unclassified/Neutral" over guessing at a violation.

4. TIMESTAMPS — Output timestamps as "YYYY-MM-DD HH:MM" (24-hour). The app \
export may use messy or inconsistent date layouts; normalize what you can. If \
only a date is known, use "YYYY-MM-DD 00:00". If no timestamp can be \
determined, output an empty string — never fabricate one.

5. EVIDENCE — Every line of the log is prefixed with its source page, e.g. \
"[Log Page 12]". Each event's evidence_reference must cite the page(s) it came \
from, e.g. "Log Page 12" or "Log Pages 12-13".

6. GRANULARITY — One row per discrete, testable event (a message thread \
awaiting reply, an exchange, an expense request, a check-in). Do not emit a \
row for every individual chat line; group lines that form one event. Do not \
duplicate events.

7. JUSTIFICATION — fact_based_justification is exactly one sentence \
analytically connecting the event to the specific order provision it satisfied \
or breached (or stating why it is untestable).

============================ COURT ORDER (RULEBOOK) ============================
{court_order}
================================================================================
"""

USER_PROMPT_TEMPLATE = """\
Below is chunk {chunk_number} of {total_chunks} of the raw coParenter log \
export. Extract every discrete, testable event and classify it per your rules. \
This chunk may begin or end mid-conversation; analyze only what is visible and \
do not speculate about content outside this chunk.

---------------------------- LOG CHUNK {chunk_number}/{total_chunks} ----------------------------
{chunk_text}
--------------------------------------------------------------------------------
"""


# --------------------------------------------------------------------------- #
# Structured output schema (Pydantic -> enforced JSON schema on the API)
# --------------------------------------------------------------------------- #

class ComplianceEvent(BaseModel):
    """One discrete, testable event extracted from the co-parenting log."""

    timestamp: str = Field(
        description='Event time formatted "YYYY-MM-DD HH:MM" (24h). '
                    'Empty string if genuinely unknown.'
    )
    event_type: str = Field(
        description="Category of event, e.g. Chat Message, Expense Request, "
                    "Custody Exchange, Geolocation Check-in, Schedule Change Request."
    )
    actor: str = Field(
        description="The parent who performed the action, exactly as identified "
                    "in the log (e.g. Father, Mother, Parent A, Parent B)."
    )
    court_order_requirement: str = Field(
        description="The specific rule, provision, or timeline from the Court "
                    "Order being tested by this event. 'None applicable' if the "
                    "order does not govern this event."
    )
    actual_event: str = Field(
        description="A cold, factual, emotion-free summary of what took place."
    )
    compliance_status: Literal[
        "Compliant",
        "Violation: Late/No-Show",
        "Violation: Response Delay",
        "Violation: Non-Payment",
        "Violation: Gatekeeping",
        "Unclassified/Neutral",
    ] = Field(description="The strict classification label.")
    fact_based_justification: str = Field(
        description="One analytical sentence connecting the event to the order "
                    "provision it satisfied or breached."
    )
    evidence_reference: str = Field(
        description='Source location in the export, e.g. "Log Page 12".'
    )


class ChunkAnalysis(BaseModel):
    """All events found in one chunk of the log."""

    events: List[ComplianceEvent] = Field(
        description="Every discrete testable event in this chunk, in "
                    "chronological order. Empty list if the chunk contains no "
                    "analyzable events."
    )


# --------------------------------------------------------------------------- #
# PDF extraction with page tracking
# --------------------------------------------------------------------------- #

@dataclass
class LogLine:
    page: int          # 1-indexed source page in the PDF
    text: str          # raw line text (already stripped)

    def tagged(self) -> str:
        return f"[Log Page {self.page}] {self.text}"


def _extract_pages_pypdf(pdf_path: Path) -> List[str]:
    from pypdf import PdfReader

    reader = PdfReader(str(pdf_path))
    pages: List[str] = []
    for page in reader.pages:
        try:
            pages.append(page.extract_text() or "")
        except Exception as exc:  # a single corrupt page must not kill the run
            print(f"  [warn] pypdf failed on page {len(pages) + 1}: {exc}")
            pages.append("")
    return pages


def _extract_pages_pdfplumber(pdf_path: Path) -> List[str]:
    import pdfplumber

    pages: List[str] = []
    with pdfplumber.open(str(pdf_path)) as pdf:
        for i, page in enumerate(pdf.pages, start=1):
            try:
                pages.append(page.extract_text() or "")
            except Exception as exc:
                print(f"  [warn] pdfplumber failed on page {i}: {exc}")
                pages.append("")
    return pages


def extract_log_lines(pdf_path: Path) -> List[LogLine]:
    """Extract text from the PDF, preferring pypdf and falling back to
    pdfplumber when pypdf produces effectively no text (common with
    oddly-encoded app exports)."""
    if not pdf_path.exists():
        raise FileNotFoundError(
            f"Co-parenting log not found: {pdf_path}\n"
            f"Place the coParenter export at that path (or pass --pdf) and re-run."
        )

    print(f"[1/4] Extracting text from {pdf_path.name} ...")
    pages: List[str] = []
    try:
        pages = _extract_pages_pypdf(pdf_path)
    except Exception as exc:
        print(f"  [warn] pypdf could not open the file ({exc}); trying pdfplumber.")

    total_chars = sum(len(p.strip()) for p in pages)
    if total_chars < 50:  # essentially empty -> try the heavier extractor
        try:
            plumber_pages = _extract_pages_pdfplumber(pdf_path)
            if sum(len(p.strip()) for p in plumber_pages) > total_chars:
                pages = plumber_pages
        except ImportError:
            print("  [warn] pdfplumber not installed; keeping pypdf output.")
        except Exception as exc:
            print(f"  [warn] pdfplumber fallback failed: {exc}")

    lines: List[LogLine] = []
    empty_pages: List[int] = []
    for page_num, page_text in enumerate(pages, start=1):
        page_lines = [ln.strip() for ln in page_text.splitlines() if ln.strip()]
        if not page_lines:
            empty_pages.append(page_num)
            continue
        for ln in page_lines:
            lines.append(LogLine(page=page_num, text=ln))

    if empty_pages:
        print(f"  [warn] {len(empty_pages)} page(s) had no extractable text "
              f"(pages: {', '.join(map(str, empty_pages[:15]))}"
              f"{', ...' if len(empty_pages) > 15 else ''}). "
              f"If these pages matter, the PDF may be scanned images needing OCR.")

    if not lines:
        raise ValueError(
            f"No extractable text found anywhere in {pdf_path.name}. "
            f"The export is likely image-based; run it through OCR first."
        )

    print(f"  Extracted {len(lines)} lines across {len(pages)} pages.")
    return lines


# --------------------------------------------------------------------------- #
# Token-safe chunking
# --------------------------------------------------------------------------- #

def build_chunks(lines: List[LogLine], char_limit: int = CHUNK_CHAR_LIMIT) -> List[str]:
    """Group page-tagged lines into chunks of ~CHUNK_TOKEN_TARGET tokens each,
    never splitting a line across chunks (except pathological single lines
    longer than an entire chunk, which are hard-sliced)."""
    chunks: List[str] = []
    current: List[str] = []
    current_len = 0

    def flush() -> None:
        nonlocal current, current_len
        if current:
            chunks.append("\n".join(current))
            current, current_len = [], 0

    for line in lines:
        tagged = line.tagged()
        # Pathological case: one "line" bigger than a whole chunk.
        if len(tagged) > char_limit:
            flush()
            for start in range(0, len(tagged), char_limit):
                chunks.append(tagged[start:start + char_limit])
            continue
        if current_len + len(tagged) + 1 > char_limit:
            flush()
        current.append(tagged)
        current_len += len(tagged) + 1

    flush()
    print(f"[2/4] Split log into {len(chunks)} token-safe chunk(s) "
          f"(~{CHUNK_TOKEN_TARGET:,} tokens each).")
    return chunks


# --------------------------------------------------------------------------- #
# Agentic compliance evaluation (LLM step)
# --------------------------------------------------------------------------- #

@dataclass
class PipelineStats:
    chunks_ok: int = 0
    chunks_failed: int = 0
    failures: List[str] = field(default_factory=list)


def load_court_order(order_path: Path) -> str:
    if not order_path.exists():
        raise FileNotFoundError(
            f"Court Order rulebook not found: {order_path}\n"
            f"Save the controlling order's text as plain UTF-8 at that path "
            f"(or pass --order) and re-run."
        )
    text = order_path.read_text(encoding="utf-8", errors="replace").strip()
    if not text:
        raise ValueError(f"{order_path} is empty — the rulebook is required.")
    return text


def make_client() -> anthropic.Anthropic:
    # The SDK resolves credentials from ANTHROPIC_API_KEY, ANTHROPIC_AUTH_TOKEN,
    # or an `ant auth login` profile. We construct lazily and translate auth
    # failures into an actionable message at call time.
    return anthropic.Anthropic()


def analyze_chunk(
    client: anthropic.Anthropic,
    system_blocks: list,
    chunk_text: str,
    chunk_number: int,
    total_chunks: int,
) -> ChunkAnalysis:
    """Send one chunk for analysis with schema-enforced structured output and
    exponential-backoff retries on transient failures."""
    user_prompt = USER_PROMPT_TEMPLATE.format(
        chunk_number=chunk_number,
        total_chunks=total_chunks,
        chunk_text=chunk_text,
    )

    last_exc: Optional[Exception] = None
    for attempt in range(MAX_API_RETRIES + 1):
        try:
            response = client.messages.parse(
                model=MODEL,
                max_tokens=MAX_OUTPUT_TOKENS,
                system=system_blocks,
                messages=[{"role": "user", "content": user_prompt}],
                output_format=ChunkAnalysis,
            )

            if response.stop_reason == "refusal":
                raise RuntimeError(
                    "the model declined this chunk (stop_reason=refusal)"
                )
            if response.stop_reason == "max_tokens":
                print(f"  [warn] chunk {chunk_number}: output hit the "
                      f"{MAX_OUTPUT_TOKENS}-token cap; results for this chunk "
                      f"may be incomplete. Consider lowering CHUNK_TOKEN_TARGET.")

            parsed = response.parsed_output
            if parsed is None:
                raise RuntimeError("model returned no parseable structured output")
            return parsed

        except anthropic.AuthenticationError:
            raise SystemExit(
                "\nERROR: Anthropic authentication failed.\n"
                "Set your API key first, e.g.:\n"
                "    export ANTHROPIC_API_KEY=sk-ant-...\n"
                "or log in once with:  ant auth login\n"
            )
        except anthropic.BadRequestError as exc:
            # Non-retryable: a malformed request will not fix itself.
            raise RuntimeError(f"API rejected the request: {exc.message}") from exc
        except (anthropic.RateLimitError,
                anthropic.APIConnectionError,
                anthropic.APIStatusError) as exc:
            status = getattr(exc, "status_code", None)
            retryable = (
                isinstance(exc, (anthropic.RateLimitError, anthropic.APIConnectionError))
                or (status is not None and status >= 500)
            )
            if not retryable:
                raise RuntimeError(f"API error on chunk {chunk_number}: {exc}") from exc
            last_exc = exc
            if attempt < MAX_API_RETRIES:
                delay = min(2 ** (attempt + 1) + random.uniform(0, 1), 60)
                print(f"  [warn] transient API error on chunk {chunk_number} "
                      f"({type(exc).__name__}); retry {attempt + 1}/"
                      f"{MAX_API_RETRIES} in {delay:.1f}s")
                time.sleep(delay)

    raise RuntimeError(
        f"chunk {chunk_number} failed after {MAX_API_RETRIES} retries: {last_exc}"
    )


def run_llm_pipeline(
    client: anthropic.Anthropic,
    court_order: str,
    chunks: List[str],
) -> tuple[List[ComplianceEvent], PipelineStats]:
    # The system prompt (instructions + rulebook) is identical for every chunk,
    # so cache it — chunk 2 onward reads the prefix at ~10% of input cost.
    system_blocks = [
        {
            "type": "text",
            "text": SYSTEM_PROMPT_TEMPLATE.format(court_order=court_order),
            "cache_control": {"type": "ephemeral"},
        }
    ]

    print(f"[3/4] Analyzing {len(chunks)} chunk(s) with {MODEL} ...")
    all_events: List[ComplianceEvent] = []
    stats = PipelineStats()

    for i, chunk_text in enumerate(chunks, start=1):
        print(f"  chunk {i}/{len(chunks)} ...", end=" ", flush=True)
        try:
            analysis = analyze_chunk(client, system_blocks, chunk_text, i, len(chunks))
            all_events.extend(analysis.events)
            stats.chunks_ok += 1
            print(f"{len(analysis.events)} event(s)")
        except Exception as exc:
            stats.chunks_failed += 1
            msg = f"chunk {i}: {exc}"
            stats.failures.append(msg)
            print(f"FAILED ({exc})")
            # A failed chunk becomes a visible placeholder row so the gap is
            # documented in the deliverable instead of silently dropped.
            pages = sorted(set(int(m) for m in re.findall(r"\[Log Page (\d+)\]", chunk_text)))
            page_ref = (f"Log Pages {pages[0]}-{pages[-1]}" if len(pages) > 1
                        else f"Log Page {pages[0]}" if pages else "Unknown")
            all_events.append(ComplianceEvent(
                timestamp="",
                event_type="Processing Gap",
                actor="N/A",
                court_order_requirement="N/A",
                actual_event=f"Automated analysis of this log segment failed "
                             f"({type(exc).__name__}). Manual review required.",
                compliance_status="Unclassified/Neutral",
                fact_based_justification="This segment could not be evaluated "
                                         "by the automated pipeline.",
                evidence_reference=page_ref,
            ))

    return all_events, stats


# --------------------------------------------------------------------------- #
# Timestamp normalization (defensive against messy app date layouts)
# --------------------------------------------------------------------------- #

_KNOWN_FORMATS = (
    "%Y-%m-%d %H:%M", "%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M:%S", "%Y-%m-%d",
    "%m/%d/%Y %I:%M %p", "%m/%d/%Y %H:%M", "%m/%d/%Y", "%m/%d/%y %I:%M %p",
    "%m/%d/%y", "%b %d, %Y %I:%M %p", "%b %d, %Y", "%B %d, %Y %I:%M %p",
    "%B %d, %Y", "%d %b %Y %H:%M", "%d %b %Y",
)


def normalize_timestamp(raw: str) -> tuple[str, Optional[datetime]]:
    """Return (display_string, sort_key). Falls back to the raw string when
    the value cannot be parsed, so no evidence is ever discarded."""
    raw = (raw or "").strip()
    if not raw:
        return "", None
    for fmt in _KNOWN_FORMATS:
        try:
            dt = datetime.strptime(raw, fmt)
            return dt.strftime("%Y-%m-%d %H:%M"), dt
        except ValueError:
            continue
    if dateutil_parser is not None:
        try:
            dt = dateutil_parser.parse(raw, fuzzy=True)
            return dt.strftime("%Y-%m-%d %H:%M"), dt
        except (ValueError, OverflowError):
            pass
    return raw, None  # keep the original text; it is still evidence


# --------------------------------------------------------------------------- #
# Spreadsheet deployment
# --------------------------------------------------------------------------- #

def build_dataframe(events: List[ComplianceEvent]) -> pd.DataFrame:
    rows = []
    for ev in events:
        display_ts, sort_key = normalize_timestamp(ev.timestamp)
        status = (ev.compliance_status if ev.compliance_status in COMPLIANCE_STATUSES
                  else "Unclassified/Neutral")
        rows.append({
            "Timestamp": display_ts,
            "Event Type": ev.event_type.strip(),
            "Actor": ev.actor.strip(),
            "Court Order Requirement": ev.court_order_requirement.strip(),
            "Actual Event": ev.actual_event.strip(),
            "Compliance Status": status,
            "Fact-Based Justification": ev.fact_based_justification.strip(),
            "Evidence Reference": ev.evidence_reference.strip(),
            "_sort": sort_key,
        })

    df = pd.DataFrame(rows, columns=OUTPUT_COLUMNS + ["_sort"])
    if not df.empty:
        # Chronological order; unparseable/blank timestamps sink to the bottom
        # in their original discovery order.
        df["_orig"] = range(len(df))
        df = df.sort_values(by=["_sort", "_orig"], na_position="last")
        df = df.drop(columns=["_sort", "_orig"]).reset_index(drop=True)
    else:
        df = df.drop(columns=["_sort"])
    return df


_COLUMN_WIDTHS = {
    "Timestamp": 17,
    "Event Type": 20,
    "Actor": 12,
    "Court Order Requirement": 42,
    "Actual Event": 55,
    "Compliance Status": 26,
    "Fact-Based Justification": 55,
    "Evidence Reference": 16,
}


def export_spreadsheet(df: pd.DataFrame, outputs_dir: Path) -> Path:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    outputs_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    xlsx_path = outputs_dir / f"compliance_report_{stamp}.xlsx"
    csv_path = outputs_dir / f"compliance_report_{stamp}.csv"

    print(f"[4/4] Writing {len(df)} event(s) to {xlsx_path.name} ...")

    with pd.ExcelWriter(xlsx_path, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="Compliance Log", index=False)
        ws = writer.sheets["Compliance Log"]

        header_font = Font(bold=True, color="FFFFFF", size=11)
        header_fill = PatternFill(start_color="1F3864", end_color="1F3864",
                                  fill_type="solid")
        violation_fill = PatternFill(start_color="FCE4E4", end_color="FCE4E4",
                                     fill_type="solid")
        compliant_fill = PatternFill(start_color="E7F4E4", end_color="E7F4E4",
                                     fill_type="solid")
        wrap = Alignment(wrap_text=True, vertical="top")

        for col_idx, col_name in enumerate(OUTPUT_COLUMNS, start=1):
            letter = get_column_letter(col_idx)
            ws.column_dimensions[letter].width = _COLUMN_WIDTHS.get(col_name, 20)
            cell = ws.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(vertical="center")

        status_col = OUTPUT_COLUMNS.index("Compliance Status") + 1
        for row_idx in range(2, ws.max_row + 1):
            status_value = ws.cell(row=row_idx, column=status_col).value or ""
            row_fill = None
            if str(status_value).startswith("Violation"):
                row_fill = violation_fill
            elif status_value == "Compliant":
                row_fill = compliant_fill
            for col_idx in range(1, len(OUTPUT_COLUMNS) + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.alignment = wrap
                if row_fill is not None:
                    cell.fill = row_fill

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    # CSV twin for lawyers/paralegals who prefer plain data.
    df.to_csv(csv_path, index=False)
    print(f"  Also wrote CSV twin: {csv_path.name}")
    return xlsx_path


# --------------------------------------------------------------------------- #
# Entry point
# --------------------------------------------------------------------------- #

def main() -> int:
    ap = argparse.ArgumentParser(
        description="Parse a coParenter log PDF, evaluate it against a Court "
                    "Order rulebook with an LLM, and export a lawyer-ready "
                    "compliance spreadsheet.",
    )
    ap.add_argument("--pdf", type=Path, default=DEFAULT_PDF_PATH,
                    help=f"coParenter export PDF (default: {DEFAULT_PDF_PATH})")
    ap.add_argument("--order", type=Path, default=DEFAULT_ORDER_PATH,
                    help=f"Court Order rulebook text file (default: {DEFAULT_ORDER_PATH})")
    ap.add_argument("--outputs", type=Path, default=OUTPUTS_DIR,
                    help=f"Output directory (default: {OUTPUTS_DIR})")
    args = ap.parse_args()

    try:
        court_order = load_court_order(args.order)
        lines = extract_log_lines(args.pdf)
        chunks = build_chunks(lines)
        client = make_client()
        events, stats = run_llm_pipeline(client, court_order, chunks)
        df = build_dataframe(events)
        xlsx_path = export_spreadsheet(df, args.outputs)
    except (FileNotFoundError, ValueError) as exc:
        print(f"\nERROR: {exc}", file=sys.stderr)
        return 1
    except KeyboardInterrupt:
        print("\nInterrupted.", file=sys.stderr)
        return 130

    print("\n" + "=" * 60)
    print("DONE")
    print(f"  Events extracted : {len(df)}")
    if not df.empty:
        for status, count in df["Compliance Status"].value_counts().items():
            print(f"    {status:<28} {count}")
    print(f"  Chunks processed : {stats.chunks_ok} ok / {stats.chunks_failed} failed")
    for failure in stats.failures:
        print(f"    [gap] {failure}")
    print(f"  Deliverable      : {xlsx_path}")
    print("=" * 60)
    return 0 if stats.chunks_failed == 0 else 2


if __name__ == "__main__":
    sys.exit(main())
